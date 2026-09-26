#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Crea el archivo de la hoja 'Citas barberia' para subir a Google Sheets.

Las columnas se extraen DIRECTAMENTE del workflow entregado, no se escriben
a mano: asi no hay errores de tecleo (acentos, espacios finales, etc).

Genera:
  - Citas barberia.xlsx   (recomendado: conserva los encabezados exactos)
  - Citas barberia.csv    (alternativa, con BOM UTF-8 para que no se rompan
                           los acentos al importarlo)

Uso: uv run --with openpyxl python crear-hoja.py
"""
import csv
import json
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

BASE = r"G:\Barberia"
ORIGEN = r"G:\Barberia\BarberiaAgenteFLUJO-1-UNCENSORED.json"
XLSX = r"G:\Barberia\Citas barberia.xlsx"
CSV = r"G:\Barberia\Citas barberia.csv"
NOMBRE_PESTANA = "Hoja 1"


def columnas_del_flujo():
    """Lee los encabezados exactos que espera el nodo de registro."""
    d = json.load(open(ORIGEN, encoding="utf-8"))
    wf = d[0] if isinstance(d, list) else d
    nodo = next(n for n in wf["nodes"]
                if n["name"] == "Registrar en hoja de citas")
    cols = nodo["parameters"]["columns"]["value"]
    return list(cols.keys())


def main():
    cols = columnas_del_flujo()
    print("Columnas extraidas del workflow (exactas):")
    for i, c in enumerate(cols, 1):
        print(f"  {i}. {c!r}")
    print()

    # --- Excel -------------------------------------------------------
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = NOMBRE_PESTANA

    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = max(16, len(cols[i - 1]) + 4)
    ws.freeze_panes = "A2"
    wb.save(XLSX)
    print(f"XLSX creado: {XLSX}")

    # --- CSV ---------------------------------------------------------
    # utf-8-sig pone el BOM: sin el, Google Sheets puede romper los acentos.
    with open(CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
    print(f"CSV  creado: {CSV}")

    print()
    print("Pestana:", NOMBRE_PESTANA)
    print("Filas: solo el encabezado (el agente agrega las citas).")
    return 0


if __name__ == "__main__":
    sys.exit(main())